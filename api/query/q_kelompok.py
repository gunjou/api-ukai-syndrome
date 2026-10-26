import io
import re
import csv
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError

from ..utils.helper import serialize_value
from ..utils.config import get_connection, get_wib


# ============================================================================ #
#                               #ANCHOR - HELPER                               #
# ============================================================================ #

def is_valid_batch(id_batch):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT id_batch FROM batch WHERE id_batch = :id_batch AND status = 1
            """), {
                "id_batch": id_batch
            }).scalar()

            return result is not None

    except SQLAlchemyError as e:
        print(f"[is_valid_batch] Error: {e}")
        return False


def is_nama_kelompok_exists(id_batch, nama_kelompok, exclude_id_kelompok=None):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            query = """
                SELECT 1
                FROM kelompok_belajar
                WHERE id_batch = :id_batch
                  AND LOWER(nama_kelompok) = LOWER(:nama_kelompok)
                  AND status = 1
            """

            params = {
                "id_batch": id_batch,
                "nama_kelompok": nama_kelompok
            }

            if exclude_id_kelompok is not None:
                query += " AND id_kelompok <> :exclude_id_kelompok"
                params["exclude_id_kelompok"] = exclude_id_kelompok

            result = conn.execute(text(query), params).fetchone()

            return bool(result)

    except SQLAlchemyError as e:
        print(f"[is_nama_kelompok_exists] Error: {e}")
        return False


def count_anggota_kelompok(id_kelompok):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT COUNT(*)
                FROM kelompok_belajar_peserta kbp
                INNER JOIN users u
                    ON u.id_user = kbp.id_peserta
                   AND u.status = 1
                WHERE kbp.id_kelompok = :id_kelompok
                  AND kbp.status = 1
            """), {
                "id_kelompok": id_kelompok
            }).scalar()

            return result or 0

    except SQLAlchemyError as e:
        print(f"[count_anggota_kelompok] Error: {e}")
        return 0


def has_upcoming_jadwal_kelompok(id_kelompok):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT 1
                FROM jadwal_kelas
                WHERE id_kelompok = :id_kelompok
                  AND status = 1
                  AND COALESCE(tanggal_reschedule, tanggal) >= :today
                LIMIT 1
            """), {
                "id_kelompok": id_kelompok,
                "today": get_wib().date()
            }).fetchone()

            return bool(result)

    except SQLAlchemyError as e:
        print(f"[has_upcoming_jadwal_kelompok] Error: {e}")
        return False



# ============================================================================ #
#                      #ANCHOR - MANAJEMEN KELOMPOK BELAJAR                    #
# ============================================================================ #

def get_all_kelompok(id_batch=None, search=None, page=1, limit=20):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            where_clause = "WHERE kb.status = 1"
            params = {}

            if id_batch is not None:
                where_clause += " AND kb.id_batch = :id_batch"
                params["id_batch"] = id_batch

            if search and search.strip():
                where_clause += """
                    AND (
                        kb.nama_kelompok ILIKE :search
                        OR b.nama_batch ILIKE :search
                    )
                """
                params["search"] = f"%{search.strip()}%"

            total = conn.execute(
                text(f"""
                    SELECT COUNT(*)
                    FROM kelompok_belajar kb
                    LEFT JOIN batch b ON b.id_batch = kb.id_batch
                    {where_clause}
                """),
                params
            ).scalar() or 0

            params["limit"] = limit
            params["offset"] = (page - 1) * limit

            result = conn.execute(
                text(f"""
                    SELECT
                        kb.id_kelompok, kb.id_batch, b.nama_batch, kb.nama_kelompok,
                        COALESCE(ang.jumlah_anggota, 0) AS jumlah_anggota,
                        kb.status, kb.created_at, kb.updated_at
                    FROM kelompok_belajar kb
                    LEFT JOIN batch b
                        ON b.id_batch = kb.id_batch
                    LEFT JOIN (
                        SELECT kbp.id_kelompok, COUNT(*) AS jumlah_anggota
                        FROM kelompok_belajar_peserta kbp
                        INNER JOIN users u
                            ON u.id_user = kbp.id_peserta
                           AND u.status = 1
                        WHERE kbp.status = 1
                        GROUP BY kbp.id_kelompok
                    ) ang
                        ON ang.id_kelompok = kb.id_kelompok
                    {where_clause}
                    ORDER BY b.nama_batch ASC, kb.nama_kelompok ASC, kb.id_kelompok ASC
                    LIMIT :limit
                    OFFSET :offset
                """),
                params
            ).mappings().fetchall()

            total_page = (total + limit - 1) // limit if total > 0 else 0

            return {
                "data": [serialize_value(row) for row in result],
                "total": total,
                "page": page,
                "limit": limit,
                "total_page": total_page
            }

    except SQLAlchemyError as e:
        print(f"[get_all_kelompok] Error: {e}")
        raise


def get_dropdown_kelompok(id_batch=None):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            query = """
                SELECT kb.id_kelompok, kb.nama_kelompok, kb.id_batch, b.nama_batch
                FROM kelompok_belajar kb
                LEFT JOIN batch b
                    ON b.id_batch = kb.id_batch
                WHERE kb.status = 1
            """

            params = {}

            if id_batch is not None:
                query += " AND kb.id_batch = :id_batch"
                params["id_batch"] = id_batch

            query += " ORDER BY kb.nama_kelompok ASC"

            result = conn.execute(text(query), params).mappings().fetchall()

            return [serialize_value(row) for row in result]

    except SQLAlchemyError as e:
        print(f"[get_dropdown_kelompok] Error: {e}")
        raise


def get_kelompok_by_id(id_kelompok):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT
                    kb.id_kelompok, kb.id_batch, b.nama_batch, kb.nama_kelompok,
                    kb.status, kb.created_at, kb.updated_at
                FROM kelompok_belajar kb
                LEFT JOIN batch b
                    ON b.id_batch = kb.id_batch
                WHERE kb.id_kelompok = :id_kelompok
                  AND kb.status = 1
            """), {
                "id_kelompok": id_kelompok
            }).mappings().fetchone()

            return serialize_value(result) if result else None

    except SQLAlchemyError as e:
        print(f"[get_kelompok_by_id] Error: {e}")
        return None


def insert_kelompok(payload):
    engine = get_connection()

    try:
        with engine.begin() as conn:
            result = conn.execute(text("""
                INSERT INTO kelompok_belajar (
                    id_batch, nama_kelompok, status, created_at, updated_at
                )
                VALUES (
                    :id_batch, :nama_kelompok, 1, :now, :now
                )
                RETURNING id_kelompok
            """), {
                **payload,
                "now": get_wib()
            }).scalar()

        if not result:
            return None

        return get_kelompok_by_id(result)

    except SQLAlchemyError as e:
        print(f"[insert_kelompok] Error: {e}")
        return None


def update_kelompok(payload):
    engine = get_connection()

    try:
        with engine.begin() as conn:
            result = conn.execute(text("""
                UPDATE kelompok_belajar
                SET
                    id_batch = :id_batch,
                    nama_kelompok = :nama_kelompok,
                    updated_at = :now
                WHERE id_kelompok = :id_kelompok
                  AND status = 1
                RETURNING id_kelompok
            """), {
                **payload,
                "now": get_wib()
            }).scalar()

        if not result:
            return None

        return get_kelompok_by_id(result)

    except SQLAlchemyError as e:
        print(f"[update_kelompok] Error: {e}")
        return None


def delete_kelompok(id_kelompok):
    engine = get_connection()

    try:
        with engine.begin() as conn:
            now = get_wib()

            result = conn.execute(text("""
                UPDATE kelompok_belajar
                SET
                    status = 0,
                    updated_at = :now
                WHERE id_kelompok = :id_kelompok
                  AND status = 1
                RETURNING id_kelompok
            """), {
                "id_kelompok": id_kelompok,
                "now": now
            }).scalar()

            if not result:
                return False

            # ==========================================
            # Nonaktifkan seluruh anggota kelompok
            # supaya peserta bisa masuk kelompok lain
            # ==========================================

            conn.execute(text("""
                UPDATE kelompok_belajar_peserta
                SET
                    status = 0,
                    updated_at = :now
                WHERE id_kelompok = :id_kelompok
                  AND status = 1
            """), {
                "id_kelompok": id_kelompok,
                "now": now
            })

            return True

    except SQLAlchemyError as e:
        print(f"[delete_kelompok] Error: {e}")
        return False



# ============================================================================ #
#                     #ANCHOR - ANGGOTA KELOMPOK BELAJAR                       #
# ============================================================================ #

def get_anggota_kelompok(id_kelompok):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT
                    kbp.id_kelompok_peserta, kbp.id_peserta,
                    u.nama, u.nickname, u.email, u.no_hp,
                    (
                        SELECT STRING_AGG(DISTINCT pk.nama_kelas, ', ' ORDER BY pk.nama_kelas)
                        FROM pesertakelas pks
                        INNER JOIN paketkelas pk
                            ON pk.id_paketkelas = pks.id_paketkelas
                           AND pk.status = 1
                        WHERE pks.id_user = kbp.id_peserta
                          AND pks.status = 1
                          AND pk.id_batch = kb.id_batch
                    ) AS nama_kelas,
                    kbp.created_at
                FROM kelompok_belajar_peserta kbp
                INNER JOIN kelompok_belajar kb
                    ON kb.id_kelompok = kbp.id_kelompok
                   AND kb.status = 1
                INNER JOIN users u
                    ON u.id_user = kbp.id_peserta
                   AND u.role = 'peserta'
                   AND u.status = 1
                WHERE kbp.id_kelompok = :id_kelompok
                  AND kbp.status = 1
                ORDER BY u.nama ASC
            """), {
                "id_kelompok": id_kelompok
            }).mappings().fetchall()

            return [serialize_value(row) for row in result]

    except SQLAlchemyError as e:
        print(f"[get_anggota_kelompok] Error: {e}")
        raise


def get_peserta_tersedia_kelompok(id_batch, search=None):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            where_clause = ""
            params = {
                "id_batch": id_batch
            }

            if search and search.strip():
                where_clause = """
                    AND (
                        u.nama ILIKE :search
                        OR u.email ILIKE :search
                        OR pk.nama_kelas ILIKE :search
                    )
                """
                params["search"] = f"%{search.strip()}%"

            # ==========================================
            # Peserta yang terdaftar di kelas pada batch
            # tsb dan belum memiliki kelompok aktif
            # di batch yang sama
            # ==========================================

            result = conn.execute(
                text(f"""
                    SELECT
                        u.id_user AS id_peserta, u.nama, u.nickname, u.email,
                        STRING_AGG(DISTINCT pk.nama_kelas, ', ' ORDER BY pk.nama_kelas) AS nama_kelas
                    FROM pesertakelas pks
                    INNER JOIN paketkelas pk
                        ON pk.id_paketkelas = pks.id_paketkelas
                       AND pk.status = 1
                       AND pk.id_batch = :id_batch
                    INNER JOIN users u
                        ON u.id_user = pks.id_user
                       AND u.role = 'peserta'
                       AND u.status = 1
                    WHERE pks.status = 1
                      AND NOT EXISTS (
                          SELECT 1
                          FROM kelompok_belajar_peserta kbp
                          INNER JOIN kelompok_belajar kb
                              ON kb.id_kelompok = kbp.id_kelompok
                             AND kb.status = 1
                          WHERE kbp.id_peserta = u.id_user
                            AND kbp.status = 1
                            AND kb.id_batch = :id_batch
                      )
                      {where_clause}
                    GROUP BY u.id_user, u.nama, u.nickname, u.email
                    ORDER BY u.nama ASC
                """),
                params
            ).mappings().fetchall()

            return [serialize_value(row) for row in result]

    except SQLAlchemyError as e:
        print(f"[get_peserta_tersedia_kelompok] Error: {e}")
        raise


def assign_peserta_kelompok(id_kelompok, id_batch, list_id_peserta):
    engine = get_connection()

    try:
        with engine.begin() as conn:

            # ==========================================
            # Lock per batch, mencegah 2 request paralel
            # memasukkan peserta yang sama ke 2 kelompok
            # berbeda dalam batch yang sama
            # ==========================================

            conn.execute(text("""
                SELECT pg_advisory_xact_lock(hashtext('kelompok_belajar'), :id_batch)
            """), {
                "id_batch": id_batch
            })

            # ==========================================
            # 1. Peserta harus terdaftar di kelas
            #    pada batch kelompok tsb
            # ==========================================

            valid_rows = conn.execute(text("""
                SELECT DISTINCT pks.id_user
                FROM pesertakelas pks
                INNER JOIN paketkelas pk
                    ON pk.id_paketkelas = pks.id_paketkelas
                   AND pk.status = 1
                   AND pk.id_batch = :id_batch
                INNER JOIN users u
                    ON u.id_user = pks.id_user
                   AND u.role = 'peserta'
                   AND u.status = 1
                WHERE pks.status = 1
                  AND pks.id_user = ANY(:list_id_peserta)
            """), {
                "id_batch": id_batch,
                "list_id_peserta": list_id_peserta
            }).scalars().all()

            valid_peserta = set(valid_rows)

            # ==========================================
            # 2. Peserta belum punya kelompok aktif
            #    di batch yang sama
            # ==========================================

            existing_rows = conn.execute(text("""
                SELECT kbp.id_peserta, kb.id_kelompok, kb.nama_kelompok
                FROM kelompok_belajar_peserta kbp
                INNER JOIN kelompok_belajar kb
                    ON kb.id_kelompok = kbp.id_kelompok
                   AND kb.status = 1
                WHERE kbp.status = 1
                  AND kb.id_batch = :id_batch
                  AND kbp.id_peserta = ANY(:list_id_peserta)
            """), {
                "id_batch": id_batch,
                "list_id_peserta": list_id_peserta
            }).mappings().fetchall()

            existing_peserta = {
                row["id_peserta"]: row
                for row in existing_rows
            }

            errors = []

            for id_peserta in list_id_peserta:
                if id_peserta not in valid_peserta:
                    errors.append({
                        "id_peserta": id_peserta,
                        "message": "Peserta tidak terdaftar pada kelas di batch kelompok ini"
                    })
                    continue

                existing = existing_peserta.get(id_peserta)

                if existing:
                    if existing["id_kelompok"] == id_kelompok:
                        message = "Peserta sudah terdaftar di kelompok ini"
                    else:
                        message = (
                            "Peserta sudah terdaftar di kelompok "
                            f"{existing['nama_kelompok']}"
                        )

                    errors.append({
                        "id_peserta": id_peserta,
                        "message": message
                    })

            if errors:
                return {
                    "status": "INVALID",
                    "errors": errors
                }

            # ==========================================
            # 3. Insert seluruh peserta
            # ==========================================

            now = get_wib()

            conn.execute(text("""
                INSERT INTO kelompok_belajar_peserta (
                    id_kelompok, id_peserta, status, created_at, updated_at
                )
                VALUES (
                    :id_kelompok, :id_peserta, 1, :now, :now
                )
            """), [
                {
                    "id_kelompok": id_kelompok,
                    "id_peserta": id_peserta,
                    "now": now
                }
                for id_peserta in list_id_peserta
            ])

            return {
                "status": "SUCCESS",
                "inserted_rows": len(list_id_peserta)
            }

    except SQLAlchemyError as e:
        print(f"[assign_peserta_kelompok] Error: {e}")
        return None


def get_kelompok_peserta_by_id(id_kelompok_peserta):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT
                    kbp.id_kelompok_peserta, kbp.id_kelompok, kb.nama_kelompok,
                    kbp.id_peserta, u.nama, u.nickname,
                    kbp.status, kbp.created_at, kbp.updated_at
                FROM kelompok_belajar_peserta kbp
                INNER JOIN kelompok_belajar kb
                    ON kb.id_kelompok = kbp.id_kelompok
                   AND kb.status = 1
                INNER JOIN users u
                    ON u.id_user = kbp.id_peserta
                WHERE kbp.id_kelompok_peserta = :id_kelompok_peserta
                  AND kbp.status = 1
            """), {
                "id_kelompok_peserta": id_kelompok_peserta
            }).mappings().fetchone()

            return serialize_value(result) if result else None

    except SQLAlchemyError as e:
        print(f"[get_kelompok_peserta_by_id] Error: {e}")
        return None


def delete_kelompok_peserta(id_kelompok_peserta):
    engine = get_connection()

    try:
        with engine.begin() as conn:
            result = conn.execute(text("""
                UPDATE kelompok_belajar_peserta
                SET
                    status = 0,
                    updated_at = :now
                WHERE id_kelompok_peserta = :id_kelompok_peserta
                  AND status = 1
                RETURNING id_kelompok_peserta
            """), {
                "id_kelompok_peserta": id_kelompok_peserta,
                "now": get_wib()
            }).scalar()

            return result is not None

    except SQLAlchemyError as e:
        print(f"[delete_kelompok_peserta] Error: {e}")
        return False



# ============================================================================ #
#                  #ANCHOR - BULK IMPORT ANGGOTA KELOMPOK (CONFIG)             #
# ============================================================================ #

BULK_ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
BULK_MAX_FILE_SIZE = 2 * 1024 * 1024  # 2 MB
BULK_MAX_ROWS = 1000
BULK_TEMPLATE_HEADERS = ["nama_kelompok", "email", "nama"]

# Header file dinormalisasi dulu (lihat normalize_header), lalu dipetakan ke kolom baku
BULK_HEADER_ALIASES = {
    "nama_kelompok": "nama_kelompok",
    "kelompok": "nama_kelompok",
    "email": "email",
    "e_mail": "email",
    "nama": "nama",
    "nama_peserta": "nama"
}

EMAIL_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")



# ============================================================================ #
#                 #ANCHOR - BULK IMPORT ANGGOTA KELOMPOK (NORMALISASI)         #
# ============================================================================ #

def normalize_text(value):
    """
    Rapikan teks untuk ditampilkan:
    - None / NaN -> ""
    - Unicode NFKC (non-breaking space, karakter full-width, dll diseragamkan)
    - Hapus BOM, trim, spasi ganda menjadi satu
    """

    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass

    value = unicodedata.normalize("NFKC", str(value)).replace("\ufeff", "")

    return re.sub(r"\s+", " ", value).strip()


def normalize_key(value):
    """Key pencocokan: normalize_text + casefold (lower yang lebih tepat untuk perbandingan)"""

    return normalize_text(value).casefold()


def normalize_header(value):
    """'Nama Kelompok' / 'NAMA-KELOMPOK' / ' nama_kelompok ' -> 'nama_kelompok'"""

    return re.sub(r"[\s\-]+", "_", normalize_key(value))


def read_bulk_kelompok_file(file, filename):
    extension = "." + filename.lower().rsplit(".", 1)[-1]

    file.seek(0)
    raw = file.read()

    if extension == ".csv":
        decoded = None

        for encoding in ("utf-8-sig", "cp1252"):
            try:
                decoded = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if decoded is None:
            raise ValueError("Encoding CSV tidak dikenali, simpan ulang file sebagai CSV UTF-8")

        # sep=None -> otomatis mendeteksi pemisah ',' atau ';' (CSV dari Excel versi Indonesia)
        # skip_blank_lines=False -> nomor baris di pesan error tetap sesuai file
        return pd.read_csv(
            io.StringIO(decoded),
            dtype=str,
            keep_default_na=False,
            sep=None,
            engine="python",
            skip_blank_lines=False
        )

    return pd.read_excel(
        io.BytesIO(raw),
        dtype=str,
        keep_default_na=False,
        sheet_name=0,
        engine="openpyxl"
    )



# ============================================================================ #
#                  #ANCHOR - BULK IMPORT ANGGOTA KELOMPOK (PROSES)             #
# ============================================================================ #

def import_bulk_kelompok(file, filename, id_batch, dry_run=False):

    # ------------------------------------------------------------------ #
    # 1. BACA FILE
    # ------------------------------------------------------------------ #

    try:
        df = read_bulk_kelompok_file(file, filename)

    except ValueError as e:
        return {"status": "FILE_INVALID", "message": str(e)}

    except Exception as e:
        print(f"[import_bulk_kelompok] Read Error: {e}")
        return {
            "status": "FILE_INVALID",
            "message": "File tidak dapat dibaca, pastikan format sesuai template"
        }

    # ------------------------------------------------------------------ #
    # 2. VALIDASI HEADER
    # ------------------------------------------------------------------ #

    column_map = {}

    for column in df.columns:
        target = BULK_HEADER_ALIASES.get(normalize_header(column))

        if target and target not in column_map:
            column_map[target] = column

    if "nama_kelompok" not in column_map:
        return {"status": "FILE_INVALID", "message": "Kolom nama_kelompok wajib ada"}

    if "email" not in column_map and "nama" not in column_map:
        return {
            "status": "FILE_INVALID",
            "message": "Minimal salah satu kolom email atau nama wajib ada"
        }

    # ------------------------------------------------------------------ #
    # 3. AMBIL BARIS (header = baris 1, data mulai baris 2)
    # ------------------------------------------------------------------ #

    rows = []

    for index, record in enumerate(df.to_dict("records")):
        row = {
            "row": index + 2,
            "nama_kelompok": normalize_text(record.get(column_map["nama_kelompok"])),
            "email": normalize_text(record.get(column_map["email"])) if "email" in column_map else "",
            "nama": normalize_text(record.get(column_map["nama"])) if "nama" in column_map else ""
        }

        if not (row["nama_kelompok"] or row["email"] or row["nama"]):
            continue

        rows.append(row)

    if not rows:
        return {"status": "FILE_INVALID", "message": "File tidak memiliki data"}

    if len(rows) > BULK_MAX_ROWS:
        return {
            "status": "FILE_INVALID",
            "message": f"Jumlah data maksimal {BULK_MAX_ROWS} baris per file"
        }

    # ------------------------------------------------------------------ #
    # 4. VALIDASI & SIMPAN (dalam satu transaction)
    # ------------------------------------------------------------------ #

    engine = get_connection()

    try:
        with engine.begin() as conn:

            # Lock per batch, sama seperti assign manual
            conn.execute(text("""
                SELECT pg_advisory_xact_lock(hashtext('kelompok_belajar'), :id_batch)
            """), {
                "id_batch": id_batch
            })

            # ==========================================
            # Kelompok aktif di batch
            # ==========================================

            kelompok_rows = conn.execute(text("""
                SELECT id_kelompok, nama_kelompok
                FROM kelompok_belajar
                WHERE id_batch = :id_batch
                  AND status = 1
            """), {
                "id_batch": id_batch
            }).mappings().fetchall()

            kelompok_map = {
                normalize_key(row["nama_kelompok"]): dict(row)
                for row in kelompok_rows
            }

            # ==========================================
            # Peserta yang terdaftar di kelas pada batch
            # (dimuat sekali, dicocokkan di Python)
            # ==========================================

            peserta_rows = conn.execute(text("""
                SELECT DISTINCT u.id_user, u.nama, u.email
                FROM pesertakelas pks
                INNER JOIN paketkelas pk
                    ON pk.id_paketkelas = pks.id_paketkelas
                   AND pk.status = 1
                   AND pk.id_batch = :id_batch
                INNER JOIN users u
                    ON u.id_user = pks.id_user
                   AND u.role = 'peserta'
                   AND u.status = 1
                WHERE pks.status = 1
            """), {
                "id_batch": id_batch
            }).mappings().fetchall()

            peserta_by_id = {}
            email_map = {}
            nama_map = {}

            for row in peserta_rows:
                peserta_by_id[row["id_user"]] = dict(row)

                email_key = normalize_key(row["email"])
                nama_key = normalize_key(row["nama"])

                if email_key:
                    email_map.setdefault(email_key, []).append(row["id_user"])

                if nama_key:
                    nama_map.setdefault(nama_key, []).append(row["id_user"])

            # ==========================================
            # Keanggotaan kelompok aktif di batch
            # ==========================================

            existing_rows = conn.execute(text("""
                SELECT kbp.id_peserta, kb.id_kelompok, kb.nama_kelompok
                FROM kelompok_belajar_peserta kbp
                INNER JOIN kelompok_belajar kb
                    ON kb.id_kelompok = kbp.id_kelompok
                   AND kb.status = 1
                WHERE kbp.status = 1
                  AND kb.id_batch = :id_batch
            """), {
                "id_batch": id_batch
            }).mappings().fetchall()

            existing_map = {
                row["id_peserta"]: dict(row)
                for row in existing_rows
            }

            # ==========================================
            # Validasi per baris (semua error dikumpulkan)
            # ==========================================

            errors = []
            skipped = []
            to_insert = []
            seen_peserta = {}

            for row in rows:
                messages = []
                id_peserta = None
                kelompok = None
                skip_message = None

                email_key = normalize_key(row["email"])
                nama_key = normalize_key(row["nama"])

                # --- kelompok ---
                if not row["nama_kelompok"]:
                    messages.append("nama_kelompok wajib diisi")
                else:
                    kelompok = kelompok_map.get(normalize_key(row["nama_kelompok"]))

                    if not kelompok:
                        messages.append(
                            f"Kelompok '{row['nama_kelompok']}' tidak ditemukan pada batch ini"
                        )

                # --- peserta ---
                if not email_key and not nama_key:
                    messages.append("Email atau nama peserta wajib diisi")

                elif email_key:
                    if not EMAIL_REGEX.match(row["email"]):
                        messages.append(f"Format email '{row['email']}' tidak valid")
                    else:
                        matches = email_map.get(email_key, [])

                        if not matches:
                            messages.append(
                                f"Peserta dengan email '{row['email']}' tidak ditemukan "
                                "pada kelas di batch ini"
                            )
                        elif len(matches) > 1:
                            messages.append(
                                f"Email '{row['email']}' terdaftar pada lebih dari satu peserta"
                            )
                        else:
                            id_peserta = matches[0]

                            if nama_key and id_peserta not in nama_map.get(nama_key, []):
                                messages.append(
                                    f"Nama '{row['nama']}' tidak sesuai dengan "
                                    f"pemilik email '{row['email']}'"
                                )
                                id_peserta = None

                else:
                    matches = nama_map.get(nama_key, [])

                    if not matches:
                        messages.append(
                            f"Peserta dengan nama '{row['nama']}' tidak ditemukan "
                            "pada kelas di batch ini"
                        )
                    elif len(matches) > 1:
                        messages.append(
                            f"Nama '{row['nama']}' cocok dengan {len(matches)} peserta, "
                            "gunakan email"
                        )
                    else:
                        id_peserta = matches[0]

                # --- duplikat di file ---
                if id_peserta is not None:
                    if id_peserta in seen_peserta:
                        messages.append(
                            f"Peserta sudah muncul di baris {seen_peserta[id_peserta]}"
                        )
                    else:
                        seen_peserta[id_peserta] = row["row"]

                # --- keanggotaan yang sudah ada ---
                if id_peserta is not None and kelompok:
                    existing = existing_map.get(id_peserta)

                    if existing and existing["id_kelompok"] != kelompok["id_kelompok"]:
                        messages.append(
                            f"Peserta sudah terdaftar di kelompok {existing['nama_kelompok']}"
                        )
                    elif existing:
                        skip_message = "Peserta sudah terdaftar di kelompok ini"

                # --- hasil baris ---
                if messages:
                    errors.append({**row, "messages": messages})

                elif skip_message:
                    skipped.append({**row, "message": skip_message})

                else:
                    peserta = peserta_by_id[id_peserta]

                    to_insert.append({
                        "row": row["row"],
                        "id_kelompok": kelompok["id_kelompok"],
                        "nama_kelompok": kelompok["nama_kelompok"],
                        "id_peserta": id_peserta,
                        "nama": peserta["nama"],
                        "email": peserta["email"]
                    })

            summary = {
                "total_rows": len(rows),
                "valid_rows": len(to_insert) + len(skipped),
                "invalid_rows": len(errors)
            }

            if errors:
                return {
                    "status": "INVALID",
                    **summary,
                    "errors": errors
                }

            per_kelompok = {}

            for item in to_insert:
                per_kelompok.setdefault(item["nama_kelompok"], 0)
                per_kelompok[item["nama_kelompok"]] += 1

            result = {
                **summary,
                "skipped_rows": len(skipped),
                "per_kelompok": [
                    {"nama_kelompok": nama, "jumlah": jumlah}
                    for nama, jumlah in per_kelompok.items()
                ],
                "skipped": skipped
            }

            # ==========================================
            # Dry run: hanya preview, tidak menyimpan
            # ==========================================

            if dry_run:
                return {
                    "status": "VALID",
                    **result,
                    "will_insert_rows": len(to_insert),
                    "preview": to_insert
                }

            if to_insert:
                now = get_wib()

                conn.execute(text("""
                    INSERT INTO kelompok_belajar_peserta (
                        id_kelompok, id_peserta, status, created_at, updated_at
                    )
                    VALUES (
                        :id_kelompok, :id_peserta, 1, :now, :now
                    )
                """), [
                    {
                        "id_kelompok": item["id_kelompok"],
                        "id_peserta": item["id_peserta"],
                        "now": now
                    }
                    for item in to_insert
                ])

            return {
                "status": "SUCCESS",
                **result,
                "inserted_rows": len(to_insert)
            }

    except SQLAlchemyError as e:
        print(f"[import_bulk_kelompok] Error: {e}")
        return None



# ============================================================================ #
#                 #ANCHOR - BULK IMPORT ANGGOTA KELOMPOK (TEMPLATE)            #
# ============================================================================ #

def build_template_kelompok(id_batch=None, file_format="xlsx"):
    daftar_kelompok = get_dropdown_kelompok(id_batch) if id_batch is not None else []

    contoh_kelompok = (
        daftar_kelompok[0]["nama_kelompok"]
        if daftar_kelompok
        else "Kelompok A"
    )

    contoh_rows = [
        [contoh_kelompok, "peserta1@email.com", "Nama Peserta Satu"],
        [contoh_kelompok, "", "Nama Peserta Dua"]
    ]

    # ------------------------------------------------------------------ #
    # CSV: header + contoh saja
    # ------------------------------------------------------------------ #

    if file_format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(BULK_TEMPLATE_HEADERS)
        writer.writerows(contoh_rows)

        buffer = io.BytesIO(output.getvalue().encode("utf-8-sig"))

        return buffer, "template_import_kelompok.csv", "text/csv"

    # ------------------------------------------------------------------ #
    # XLSX
    # ------------------------------------------------------------------ #

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")

    wb = Workbook()

    # Sheet 1: Import (harus sheet pertama, karena import membaca sheet_name=0)
    ws = wb.active
    ws.title = "Import"
    ws.append(BULK_TEMPLATE_HEADERS)

    for row in contoh_rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 35
    ws.freeze_panes = "A2"

    # Sheet 2: Daftar Kelompok + dropdown (jika id_batch dikirim)
    if daftar_kelompok:
        ws_kelompok = wb.create_sheet("Daftar Kelompok")
        ws_kelompok.append(["nama_kelompok"])
        ws_kelompok["A1"].font = header_font
        ws_kelompok["A1"].fill = header_fill
        ws_kelompok.column_dimensions["A"].width = 35

        for item in daftar_kelompok:
            ws_kelompok.append([item["nama_kelompok"]])

        last_row = len(daftar_kelompok) + 1

        dv = DataValidation(
            type="list",
            formula1=f"='Daftar Kelompok'!$A$2:$A${last_row}",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Kelompok tidak valid",
            error="Pilih nama kelompok dari daftar"
        )

        ws.add_data_validation(dv)
        dv.add(f"A2:A{BULK_MAX_ROWS + 1}")

    # Sheet 3: Petunjuk
    ws_petunjuk = wb.create_sheet("Petunjuk")
    ws_petunjuk.column_dimensions["A"].width = 100

    petunjuk = [
        "PETUNJUK IMPORT ANGGOTA KELOMPOK",
        "",
        "1. Isi data pada sheet 'Import'. Jangan mengubah nama kolom pada baris pertama.",
        "2. nama_kelompok wajib diisi dan harus sama dengan kelompok yang sudah terdaftar pada batch.",
        "3. Isi email peserta (disarankan). Kolom nama hanya dipakai jika email kosong.",
        "4. Jika hanya mengisi nama, nama harus unik pada batch. Jika ada nama yang sama, wajib gunakan email.",
        "5. Jika email dan nama diisi bersamaan, keduanya harus milik peserta yang sama.",
        "6. Peserta harus terdaftar di kelas pada batch tersebut.",
        "7. Satu peserta hanya boleh muncul satu kali dan hanya boleh berada di satu kelompok per batch.",
        "8. Peserta yang sudah terdaftar di kelompok yang sama akan dilewati (tidak dianggap error).",
        "9. Huruf besar/kecil dan spasi berlebih tidak berpengaruh pada pencocokan.",
        f"10. Maksimal {BULK_MAX_ROWS} baris per file. Jika ada satu baris error, tidak ada data yang disimpan.",
        "11. Hapus baris contoh sebelum mengunggah file."
    ]

    for line in petunjuk:
        ws_petunjuk.append([line])

    ws_petunjuk["A1"].font = Font(bold=True, size=12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return (
        buffer,
        "template_import_kelompok.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )