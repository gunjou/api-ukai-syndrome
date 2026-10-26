import io
import os
import calendar
from datetime import date, datetime, time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError


from ..utils.config import get_connection, get_wib


# ============================================================================ #
#                                #ANCHOR - CONFIG                              #
# ============================================================================ #

# Tarif honor mentor per 1 jam mengajar (Rupiah)
PAYROLL_TARIF_PER_JAM = int(os.getenv("PAYROLL_TARIF_PER_JAM", 100000))

STATUS_SESI_PENUH = "PENUH"
STATUS_SESI_SEBAGIAN = "SEBAGIAN"
STATUS_SESI_BELUM_CHECK_OUT = "BELUM_CHECK_OUT"
STATUS_SESI_TIDAK_HADIR = "TIDAK_HADIR"

STATUS_SESI_COUNTER = {
    STATUS_SESI_PENUH: "sesi_penuh",
    STATUS_SESI_SEBAGIAN: "sesi_sebagian",
    STATUS_SESI_BELUM_CHECK_OUT: "sesi_belum_check_out",
    STATUS_SESI_TIDAK_HADIR: "sesi_tidak_hadir"
}

NAMA_BULAN = [
    "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember"
]



# ============================================================================ #
#                               #ANCHOR - HELPER                               #
# ============================================================================ #

def get_periode_bulan(tahun, bulan):
    """Tanggal 1 s/d tanggal terakhir bulan (28/29/30/31)"""

    last_day = calendar.monthrange(tahun, bulan)[1]

    return date(tahun, bulan, 1), date(tahun, bulan, last_day)


def hitung_honor(menit):
    """Honor = menit x tarif / 60, dibulatkan ke rupiah terdekat (tanpa float)"""

    return (int(menit) * PAYROLL_TARIF_PER_JAM + 30) // 60


def format_durasi(menit):
    """165 -> '2 jam 45 menit', 180 -> '3 jam', 0 -> '0 menit'"""

    if menit is None:
        return None

    jam, sisa = divmod(int(menit), 60)

    if jam and sisa:
        return f"{jam} jam {sisa} menit"

    if jam:
        return f"{jam} jam"

    return f"{sisa} menit"


def format_value(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, time):
        return value.strftime("%H:%M:%S")

    return value


def get_mentor_payroll(id_mentor):
    """Data mentor untuk payroll (mentor nonaktif tetap bisa dilihat riwayatnya)"""

    engine = get_connection()

    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT id_user AS id_mentor, nama AS nama_mentor, nickname AS nickname_mentor, email, status
                FROM users
                WHERE id_user = :id_mentor
                  AND role = 'mentor'
            """), {
                "id_mentor": id_mentor
            }).mappings().fetchone()

            return dict(result) if result else None

    except SQLAlchemyError as e:
        print(f"[get_mentor_payroll] Error: {e}")
        return None



# ============================================================================ #
#                     #ANCHOR - SESI PAYROLL (SUMBER DATA UTAMA)               #
# ============================================================================ #
# Semua endpoint payroll (admin, mentor, export) memakai fungsi ini,
# sehingga angka di semua tampilan dijamin sama.
#
# Aturan:
# - 1 sesi = jadwal dengan mentor + tanggal + jam mulai + jam selesai yang sama
#   (sama dengan /jadwal/group). Banyak kelas dalam 1 sesi dibayar 1 kali.
# - durasi_dibayar = MIN(durasi_aktual, durasi_jadwal)
# - Belum check-out -> dibayar 0, ditandai BELUM_CHECK_OUT (koreksi manual admin)
# - Jadwal yang sudah lewat tanpa check-in -> TIDAK_HADIR
# - Jadwal yang belum terlaksana tidak ditampilkan

def _hitung_sesi(sesi, now):
    jadwal_mulai = datetime.combine(sesi["tanggal"], sesi["waktu_mulai"])
    jadwal_selesai = datetime.combine(sesi["tanggal"], sesi["waktu_selesai"])

    durasi_jadwal = max(0, int((jadwal_selesai - jadwal_mulai).total_seconds() // 60))

    check_in_at = sesi["check_in_at"]
    check_out_at = sesi["check_out_at"]

    durasi_aktual = None
    durasi_dibayar = 0

    if check_in_at is None:

        # Jadwal belum selesai dan belum ada check-in -> belum terlaksana
        if now < jadwal_selesai:
            return None

        status_sesi = STATUS_SESI_TIDAK_HADIR

    elif check_out_at is None:
        status_sesi = STATUS_SESI_BELUM_CHECK_OUT

    else:
        durasi_aktual = max(0, int((check_out_at - check_in_at).total_seconds() // 60))
        durasi_dibayar = min(durasi_aktual, durasi_jadwal)

        status_sesi = (
            STATUS_SESI_PENUH
            if durasi_dibayar >= durasi_jadwal
            else STATUS_SESI_SEBAGIAN
        )

    return {
        "id_sesi": min(kelas["id_jadwal"] for kelas in sesi["kelas"]),
        "id_mentor": sesi["id_mentor"],
        "nama_mentor": sesi["nama_mentor"],
        "nickname_mentor": sesi["nickname_mentor"],

        "tanggal": format_value(sesi["tanggal"]),
        "waktu_mulai": format_value(sesi["waktu_mulai"]),
        "waktu_selesai": format_value(sesi["waktu_selesai"]),
        "type_pertemuan": ", ".join(sorted(sesi["type_pertemuan"])),

        "jumlah_kelas": len(sesi["kelas"]),
        "kelas": sesi["kelas"],

        "check_in_at": format_value(check_in_at),
        "check_out_at": format_value(check_out_at),

        "durasi_jadwal_menit": durasi_jadwal,
        "durasi_aktual_menit": durasi_aktual,
        "durasi_dibayar_menit": durasi_dibayar,
        "durasi_dibayar_label": format_durasi(durasi_dibayar),

        "honor": hitung_honor(durasi_dibayar),
        "status_sesi": status_sesi
    }


def get_sesi_payroll(start_date, end_date, id_mentor=None):
    engine = get_connection()

    try:
        with engine.connect() as conn:
            where_clause = ""
            params = {
                "start_date": start_date,
                "end_date": end_date
            }

            if id_mentor is not None:
                where_clause = " AND COALESCE(am.id_mentor, j.id_mentor) = :id_mentor"
                params["id_mentor"] = id_mentor

            # ==========================================
            # Mentor yang dibayar = mentor yang check-in
            # (absensi_mentor), fallback ke mentor jadwal
            # jika belum ada absensi
            # ==========================================

            rows = conn.execute(
                text(f"""
                    SELECT
                        j.id_jadwal,
                        COALESCE(am.id_mentor, j.id_mentor) AS id_mentor,
                        u.nama AS nama_mentor,
                        u.nickname AS nickname_mentor,

                        j.id_paketkelas, pk.nama_kelas,
                        j.id_kelompok, kb.nama_kelompok,
                        j.topik, j.type_pertemuan,

                        COALESCE(j.tanggal_reschedule, j.tanggal) AS tanggal_efektif,
                        COALESCE(j.waktu_mulai_reschedule, j.waktu_mulai) AS waktu_mulai_efektif,
                        COALESCE(j.waktu_selesai_reschedule, j.waktu_selesai) AS waktu_selesai_efektif,

                        am.id_absensi_mentor, am.check_in_at, am.check_out_at

                    FROM jadwal_kelas j
                    LEFT JOIN absensi_mentor am
                        ON am.id_jadwal = j.id_jadwal
                       AND am.status = 1
                    LEFT JOIN users u
                        ON u.id_user = COALESCE(am.id_mentor, j.id_mentor)
                    LEFT JOIN paketkelas pk
                        ON pk.id_paketkelas = j.id_paketkelas
                    LEFT JOIN kelompok_belajar kb
                        ON kb.id_kelompok = j.id_kelompok
                    WHERE j.status = 1
                      AND COALESCE(j.tanggal_reschedule, j.tanggal) BETWEEN :start_date AND :end_date
                      {where_clause}
                    ORDER BY
                        tanggal_efektif ASC,
                        waktu_mulai_efektif ASC,
                        waktu_selesai_efektif ASC,
                        j.id_jadwal ASC
                """),
                params
            ).mappings().fetchall()

            # ==========================================
            # GROUP BY MENTOR + TANGGAL + JAM -> 1 SESI
            # ==========================================

            grouped = {}

            for row in rows:
                group_key = (
                    row["id_mentor"],
                    row["tanggal_efektif"],
                    row["waktu_mulai_efektif"],
                    row["waktu_selesai_efektif"]
                )

                if group_key not in grouped:
                    grouped[group_key] = {
                        "id_mentor": row["id_mentor"],
                        "nama_mentor": row["nama_mentor"],
                        "nickname_mentor": row["nickname_mentor"],
                        "tanggal": row["tanggal_efektif"],
                        "waktu_mulai": row["waktu_mulai_efektif"],
                        "waktu_selesai": row["waktu_selesai_efektif"],
                        "type_pertemuan": set(),
                        "kelas": [],
                        "check_in_at": None,
                        "check_out_at": None
                    }

                sesi = grouped[group_key]

                sesi["type_pertemuan"].add(row["type_pertemuan"])

                sesi["kelas"].append({
                    "id_jadwal": row["id_jadwal"],
                    "target_type": "KELOMPOK" if row["id_kelompok"] is not None else "KELAS",
                    "id_paketkelas": row["id_paketkelas"],
                    "nama_kelas": row["nama_kelas"],
                    "id_kelompok": row["id_kelompok"],
                    "nama_kelompok": row["nama_kelompok"],
                    "topik": row["topik"],
                    "check_in_at": format_value(row["check_in_at"]),
                    "check_out_at": format_value(row["check_out_at"])
                })

                # ==========================================
                # Durasi aktual sesi:
                # check-in paling awal s/d check-out paling akhir
                # dari seluruh kelas di sesi tsb
                # ==========================================

                if row["check_in_at"] is not None and (
                    sesi["check_in_at"] is None
                    or row["check_in_at"] < sesi["check_in_at"]
                ):
                    sesi["check_in_at"] = row["check_in_at"]

                if row["check_out_at"] is not None and (
                    sesi["check_out_at"] is None
                    or row["check_out_at"] > sesi["check_out_at"]
                ):
                    sesi["check_out_at"] = row["check_out_at"]

            now = get_wib().replace(tzinfo=None)
            result = []

            for sesi in grouped.values():
                item = _hitung_sesi(sesi, now)

                if item:
                    result.append(item)

            return result

    except SQLAlchemyError as e:
        print(f"[get_sesi_payroll] Error: {e}")
        raise



# ============================================================================ #
#                         #ANCHOR - RINGKASAN PAYROLL                          #
# ============================================================================ #

def _ringkasan_kosong():
    return {
        "jumlah_sesi": 0,
        "jumlah_kelas": 0,
        "sesi_penuh": 0,
        "sesi_sebagian": 0,
        "sesi_belum_check_out": 0,
        "sesi_tidak_hadir": 0,
        "total_menit_jadwal": 0,
        "total_menit_aktual": 0,
        "total_menit_dibayar": 0
    }


def _tambah_ringkasan(ringkasan, sesi):
    ringkasan["jumlah_sesi"] += 1
    ringkasan["jumlah_kelas"] += sesi["jumlah_kelas"]
    ringkasan[STATUS_SESI_COUNTER[sesi["status_sesi"]]] += 1
    ringkasan["total_menit_jadwal"] += sesi["durasi_jadwal_menit"]
    ringkasan["total_menit_aktual"] += sesi["durasi_aktual_menit"] or 0
    ringkasan["total_menit_dibayar"] += sesi["durasi_dibayar_menit"]


def _finalisasi_ringkasan(ringkasan):
    ringkasan["total_jam_jadwal_label"] = format_durasi(ringkasan["total_menit_jadwal"])
    ringkasan["total_jam_dibayar_label"] = format_durasi(ringkasan["total_menit_dibayar"])

    # Honor total dihitung dari total menit (bukan jumlah honor per sesi)
    # supaya tidak ada selisih pembulatan yang menumpuk
    ringkasan["honor"] = hitung_honor(ringkasan["total_menit_dibayar"])

    return ringkasan


def build_ringkasan(sesi_list):
    ringkasan = _ringkasan_kosong()

    for sesi in sesi_list:
        _tambah_ringkasan(ringkasan, sesi)

    return _finalisasi_ringkasan(ringkasan)


def build_rekap_mentor(sesi_list):
    rekap = {}

    for sesi in sesi_list:
        if sesi["id_mentor"] not in rekap:
            rekap[sesi["id_mentor"]] = {
                "id_mentor": sesi["id_mentor"],
                "nama_mentor": sesi["nama_mentor"],
                "nickname_mentor": sesi["nickname_mentor"],
                **_ringkasan_kosong()
            }

        _tambah_ringkasan(rekap[sesi["id_mentor"]], sesi)

    return sorted(
        [_finalisasi_ringkasan(item) for item in rekap.values()],
        key=lambda item: (item["nama_mentor"] or "").casefold()
    )



# ============================================================================ #
#                     #ANCHOR - PAYROLL BY ADMIN (SEMUA MENTOR)                #
# ============================================================================ #

def get_rekap_payroll_bulanan(start_date, end_date, search=None):
    sesi_list = get_sesi_payroll(start_date, end_date)
    rekap = build_rekap_mentor(sesi_list)

    if search and search.strip():
        keyword = search.strip().casefold()

        rekap = [
            item for item in rekap
            if keyword in (item["nama_mentor"] or "").casefold()
            or keyword in (item["nickname_mentor"] or "").casefold()
        ]

    total_menit_dibayar = sum(item["total_menit_dibayar"] for item in rekap)

    return {
        "data": rekap,
        "total": len(rekap),
        "total_menit_dibayar": total_menit_dibayar,
        "total_jam_dibayar_label": format_durasi(total_menit_dibayar),
        "total_honor": sum(item["honor"] for item in rekap)
    }



# ============================================================================ #
#                     #ANCHOR - PAYROLL BY MENTOR (DETAIL)                     #
# ============================================================================ #

def get_detail_payroll_mentor(id_mentor, start_date, end_date):
    sesi_list = get_sesi_payroll(start_date, end_date, id_mentor=id_mentor)

    rincian = [
        {
            key: value
            for key, value in sesi.items()
            if key not in ("id_mentor", "nama_mentor", "nickname_mentor")
        }
        for sesi in sesi_list
    ]

    return {
        "ringkasan": build_ringkasan(sesi_list),
        "rincian": rincian
    }


def get_riwayat_payroll_mentor(id_mentor, tahun):
    today = get_wib().date()

    if tahun > today.year:
        return []

    bulan_terakhir = today.month if tahun == today.year else 12

    start_date, _ = get_periode_bulan(tahun, 1)
    _, end_date = get_periode_bulan(tahun, bulan_terakhir)

    # 1x query untuk satu tahun, lalu dipecah per bulan
    sesi_list = get_sesi_payroll(start_date, end_date, id_mentor=id_mentor)

    per_bulan = {}

    for sesi in sesi_list:
        per_bulan.setdefault(sesi["tanggal"][:7], []).append(sesi)

    result = []

    for bulan in range(bulan_terakhir, 0, -1):
        key = f"{tahun:04d}-{bulan:02d}"

        result.append({
            "bulan": key,
            "nama_bulan": f"{NAMA_BULAN[bulan]} {tahun}",
            **build_ringkasan(per_bulan.get(key, []))
        })

    return result



# ============================================================================ #
#                        #ANCHOR - EXPORT PAYROLL (XLSX)                       #
# ============================================================================ #

def _label_kelas(sesi):
    return ", ".join(
        kelas["nama_kelas"] or kelas["nama_kelompok"] or "-"
        for kelas in sesi["kelas"]
    )


def _style_header(ws, row_number):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")

    for cell in ws[row_number]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_export_payroll(start_date, end_date, bulan):
    sesi_list = get_sesi_payroll(start_date, end_date)
    rekap = build_rekap_mentor(sesi_list)

    tahun_int, bulan_int = int(bulan[:4]), int(bulan[5:7])
    judul_periode = f"{NAMA_BULAN[bulan_int]} {tahun_int}"
    rupiah_format = '"Rp" #,##0'

    wb = Workbook()

    # ------------------------------------------------------------------ #
    # SHEET 1: REKAP PER MENTOR
    # ------------------------------------------------------------------ #

    ws = wb.active
    ws.title = "Rekap"

    ws.append([f"Rekap Payroll Mentor - {judul_periode}"])
    ws.append([f"Periode: {start_date.isoformat()} s/d {end_date.isoformat()}"])
    ws.append([f"Tarif per jam: Rp {PAYROLL_TARIF_PER_JAM:,}".replace(",", ".")])
    ws.append([])
    ws["A1"].font = Font(bold=True, size=13)

    ws.append([
        "No", "Nama Mentor", "Nickname", "Jumlah Sesi", "Jumlah Kelas",
        "Penuh", "Sebagian", "Belum Check-out", "Tidak Hadir",
        "Total Jam Jadwal", "Total Jam Dibayar", "Total Menit Dibayar", "Honor"
    ])
    header_row = ws.max_row
    _style_header(ws, header_row)

    for index, item in enumerate(rekap, start=1):
        ws.append([
            index, item["nama_mentor"], item["nickname_mentor"],
            item["jumlah_sesi"], item["jumlah_kelas"],
            item["sesi_penuh"], item["sesi_sebagian"],
            item["sesi_belum_check_out"], item["sesi_tidak_hadir"],
            item["total_jam_jadwal_label"], item["total_jam_dibayar_label"],
            item["total_menit_dibayar"], item["honor"]
        ])
        ws.cell(row=ws.max_row, column=13).number_format = rupiah_format

    total_menit = sum(item["total_menit_dibayar"] for item in rekap)

    ws.append([
        "", "TOTAL", "", sum(item["jumlah_sesi"] for item in rekap),
        sum(item["jumlah_kelas"] for item in rekap), "", "", "", "",
        "", format_durasi(total_menit), total_menit,
        sum(item["honor"] for item in rekap)
    ])

    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    ws.cell(row=ws.max_row, column=13).number_format = rupiah_format

    widths = [6, 32, 18, 12, 12, 10, 10, 15, 12, 18, 18, 18, 18]

    for column_index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=column_index).column_letter].width = width

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # ------------------------------------------------------------------ #
    # SHEET 2: RINCIAN PER SESI
    # ------------------------------------------------------------------ #

    ws_rincian = wb.create_sheet("Rincian")
    ws_rincian.append([
        "Nama Mentor", "Tanggal", "Jam Mulai", "Jam Selesai", "Kelas / Kelompok",
        "Jumlah Kelas", "Tipe", "Check-in", "Check-out",
        "Durasi Jadwal (menit)", "Durasi Aktual (menit)", "Durasi Dibayar (menit)",
        "Honor", "Status"
    ])
    _style_header(ws_rincian, 1)

    sesi_sorted = sorted(
        sesi_list,
        key=lambda sesi: (
            (sesi["nama_mentor"] or "").casefold(),
            sesi["tanggal"],
            sesi["waktu_mulai"]
        )
    )

    for sesi in sesi_sorted:
        ws_rincian.append([
            sesi["nama_mentor"], sesi["tanggal"], sesi["waktu_mulai"], sesi["waktu_selesai"],
            _label_kelas(sesi), sesi["jumlah_kelas"], sesi["type_pertemuan"],
            sesi["check_in_at"], sesi["check_out_at"],
            sesi["durasi_jadwal_menit"], sesi["durasi_aktual_menit"], sesi["durasi_dibayar_menit"],
            sesi["honor"], sesi["status_sesi"]
        ])
        ws_rincian.cell(row=ws_rincian.max_row, column=13).number_format = rupiah_format

    widths = [28, 12, 11, 11, 40, 12, 16, 20, 20, 16, 16, 16, 16, 18]

    for column_index, width in enumerate(widths, start=1):
        ws_rincian.column_dimensions[ws_rincian.cell(row=1, column=column_index).column_letter].width = width

    ws_rincian.freeze_panes = "A2"

    if ws_rincian.max_row > 1:
        ws_rincian.auto_filter.ref = ws_rincian.dimensions

    # ------------------------------------------------------------------ #
    # SHEET 3: KETERANGAN
    # ------------------------------------------------------------------ #

    ws_info = wb.create_sheet("Keterangan")
    ws_info.column_dimensions["A"].width = 110

    keterangan = [
        "KETERANGAN PERHITUNGAN PAYROLL",
        "",
        "1. Satu sesi = jadwal dengan mentor, tanggal, jam mulai, dan jam selesai yang sama. Banyak kelas dalam satu sesi dibayar satu kali.",
        "2. Durasi dibayar = durasi aktual (check-in s/d check-out), maksimal sebesar durasi jadwal.",
        "3. Tanggal dan jam mengikuti jadwal efektif (sudah memperhitungkan reschedule).",
        f"4. Honor = durasi dibayar (menit) x Rp {PAYROLL_TARIF_PER_JAM:,} / 60.".replace(",", "."),
        "5. Honor total per mentor dihitung dari total menit dibayar, sehingga bisa berbeda beberapa rupiah dari jumlah honor per sesi karena pembulatan.",
        "6. BELUM_CHECK_OUT: mentor belum check-out, dibayar 0 sampai dikoreksi admin.",
        "7. TIDAK_HADIR: jadwal sudah lewat tanpa check-in, dibayar 0.",
        "8. Rekap dihitung langsung dari data absensi saat file ini diunduh."
    ]

    for line in keterangan:
        ws_info.append([line])

    ws_info["A1"].font = Font(bold=True, size=12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return (
        buffer,
        f"rekap_payroll_mentor_{bulan}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )